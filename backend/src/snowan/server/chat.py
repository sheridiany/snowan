import base64
import io
import json
from collections.abc import AsyncIterator
from typing import Any

from fastapi import APIRouter
from fastapi.responses import StreamingResponse
from pydantic import BaseModel
from pydantic_ai import (
    Agent,
    BinaryContent,
    DeferredToolRequests,
    DeferredToolResults,
    FunctionToolCallEvent,
    FunctionToolResultEvent,
    PartDeltaEvent,
    TextPartDelta,
    ToolDenied,
)
from pydantic_ai.messages import ModelMessage

from ..agent.build import build_agent
from ..agent.sessions import delete_history, load_history, save_history

router = APIRouter()


class Attachment(BaseModel):
    name: str
    mime: str
    data: str  # base64


class ChatRequest(BaseModel):
    message: str
    session_id: str = "default"
    attachments: list[Attachment] = []


def _extract_text(name: str, mime: str, raw: bytes) -> str | None:
    """Extract readable text from office/pdf documents. None if not a known doc."""
    ext = name.lower().rsplit(".", 1)[-1] if "." in name else ""
    try:
        if ext in ("xlsx", "xlsm") or "spreadsheetml" in mime:
            from openpyxl import load_workbook

            wb = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
            out: list[str] = []
            for ws in wb.worksheets:
                out.append(f"## {ws.title}")
                for row in ws.iter_rows(values_only=True):
                    cells = [str(c) for c in row if c is not None]
                    if cells:
                        out.append(" | ".join(cells))
            return "\n".join(out)
        if ext == "docx" or "wordprocessingml" in mime:
            import docx

            d = docx.Document(io.BytesIO(raw))
            return "\n".join(p.text for p in d.paragraphs if p.text.strip())
        if ext == "pdf" or mime == "application/pdf":
            from pypdf import PdfReader

            reader = PdfReader(io.BytesIO(raw))
            return "\n\n".join((pg.extract_text() or "") for pg in reader.pages)
        if ext == "pptx" or "presentationml" in mime:
            from pptx import Presentation

            prs = Presentation(io.BytesIO(raw))
            out2: list[str] = []
            for i, slide in enumerate(prs.slides, 1):
                out2.append(f"## Slide {i}")
                for shape in slide.shapes:
                    if shape.has_text_frame and shape.text_frame.text.strip():
                        out2.append(shape.text_frame.text)
            return "\n".join(out2)
    except Exception:  # noqa: BLE001 — an unparseable doc just falls through
        return None
    return None


def _build_prompt(message: str, attachments: list[Attachment]):
    """Images -> vision content; documents -> extracted text; text files inlined."""
    if not attachments:
        return message
    parts: list[Any] = [message] if message else []
    for a in attachments:
        try:
            raw = base64.b64decode(a.data)
        except Exception:  # noqa: BLE001
            continue
        if a.mime.startswith("image/"):
            parts.append(BinaryContent(data=raw, media_type=a.mime))
            continue
        text = _extract_text(a.name, a.mime, raw)
        if text is None:
            try:
                text = raw.decode("utf-8")
            except UnicodeDecodeError:
                text = None
        if text is not None:
            parts.append(f"\n\n[附件 {a.name}]\n```\n{text[:50000]}\n```")
        else:
            parts.append(f"\n\n[附件 {a.name}:无法解析的二进制文件]")
    return parts


class ApproveRequest(BaseModel):
    session_id: str = "default"
    decisions: dict[str, bool]


def _sse(event: dict[str, Any]) -> str:
    return f"data: {json.dumps(event)}\n\n"


def _jsonable(value: Any) -> Any:
    if isinstance(value, (str, int, float, bool, type(None), dict, list)):
        return value
    return str(value)


def _approval_event(requests: DeferredToolRequests) -> dict[str, Any]:
    return {
        "type": "approval_required",
        "calls": [
            {
                "id": call.tool_call_id,
                "name": call.tool_name,
                "args": _jsonable(call.args_as_dict()),
            }
            for call in requests.approvals
        ],
    }


async def _stream_run(run: Any) -> AsyncIterator[str]:
    """Stream one agent run, emitting delta/tool_call/tool_result/approval_required events."""
    async for node in run:
        if Agent.is_model_request_node(node):
            async with node.stream(run.ctx) as request_stream:
                async for ev in request_stream:
                    if isinstance(ev, PartDeltaEvent) and isinstance(ev.delta, TextPartDelta):
                        yield _sse({"type": "delta", "text": ev.delta.content_delta})
        elif Agent.is_call_tools_node(node):
            async with node.stream(run.ctx) as handle_stream:
                async for ev in handle_stream:
                    if isinstance(ev, FunctionToolCallEvent):
                        part = ev.part
                        args = part.args
                        if isinstance(args, str):
                            try:
                                args = json.loads(args)
                            except json.JSONDecodeError:
                                pass
                        yield _sse({
                            "type": "tool_call",
                            "id": part.tool_call_id,
                            "name": part.tool_name,
                            "args": _jsonable(args),
                        })
                    elif isinstance(ev, FunctionToolResultEvent):
                        result = ev.result
                        yield _sse({
                            "type": "tool_result",
                            "id": result.tool_call_id,
                            "name": result.tool_name,
                            "result": str(result.content),
                        })

    output = run.result.output
    if isinstance(output, DeferredToolRequests) and output.approvals:
        yield _sse(_approval_event(output))
    yield _sse({"type": "done"})


async def _run_new(prompt: Any, history: list[ModelMessage], session_id: str) -> AsyncIterator[str]:
    # Build per-request so a model/key change saved in Settings takes effect at once.
    agent = build_agent()
    async with agent.iter(prompt, message_history=history) as run:
        async for chunk in _stream_run(run):
            yield chunk
        save_history(session_id, run.result.all_messages())


async def _run_resume(
    history: list[ModelMessage],
    results: DeferredToolResults,
    session_id: str,
) -> AsyncIterator[str]:
    agent = build_agent()
    async with agent.iter(message_history=history, deferred_tool_results=results) as run:
        async for chunk in _stream_run(run):
            yield chunk
        save_history(session_id, run.result.all_messages())


@router.post("/api/chat/stream")
async def chat_stream(req: ChatRequest) -> StreamingResponse:
    history = load_history(req.session_id)
    prompt = _build_prompt(req.message, req.attachments)
    return StreamingResponse(
        _run_new(prompt, history, req.session_id),
        media_type="text/event-stream",
    )


@router.delete("/api/sessions/{session_id}")
def delete_session(session_id: str) -> dict:
    delete_history(session_id)
    return {"ok": True}


@router.post("/api/chat/approve")
async def chat_approve(req: ApproveRequest) -> StreamingResponse:
    history = load_history(req.session_id)
    results = DeferredToolResults()
    for tool_call_id, approved in req.decisions.items():
        results.approvals[tool_call_id] = (
            True if approved else ToolDenied("The user denied this tool call.")
        )
    return StreamingResponse(
        _run_resume(history, results, req.session_id),
        media_type="text/event-stream",
    )
