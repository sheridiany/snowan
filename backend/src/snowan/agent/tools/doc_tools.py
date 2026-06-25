"""In-process spreadsheet/document tools for the agent: READ tabular files completely
(read_table — every sheet, every cell's value AND formula, merged ranges) and PRODUCE
Word / PDF / HTML / Markdown / Excel / PowerPoint — all with the bundled office libs, so
it works the same in dev and the packaged app (no shelling out to a system `python`, no
`pip install`). read_table is read-only; each create_* writes a workspace file and the
chat layer surfaces it as a downloadable artifact card."""
from pydantic import BaseModel

from ...export import _filename, md_to_docx, md_to_html_doc, md_to_pdf, outline_to_pptx, rows_to_xlsx
from ...workspace import resolve_in_workspace

_MAX_CELL = 200  # truncate any single cell's text in the dump


def _cell(v) -> str:
    s = "" if v is None else str(v)
    return s[:_MAX_CELL] + "…" if len(s) > _MAX_CELL else s


def read_table(path: str, sheet: str | None = None, max_rows: int = 500) -> str:
    """完整读取工作区里的表格(.xlsx/.xlsm/.csv):列出每个 sheet 的维度、合并区、逐行的单元格值,以及
    带公式单元格的公式原文。分析/对比/按模板重排前,**先用它把上传的原件读全**(含公式、多个 sheet),
    不要用临时 shell 脚本或 pip 装包去读。path 用附件块里给的 uploads/… 路径;sheet 只读某一页;
    max_rows 限制每页行数(默认 500)。"""
    try:
        p = resolve_in_workspace(path)
    except ValueError:
        return f"error: 路径越界:{path}"
    if not p.exists():
        return f"error: 文件不存在:{path}(用附件块里标注的 uploads/… 路径)"
    ext = p.suffix.lower()
    if ext == ".csv":
        try:
            import csv as _csv

            with p.open("r", encoding="utf-8-sig", newline="") as f:
                rows = list(_csv.reader(f))
        except Exception as e:  # noqa: BLE001
            return f"error: 读取 CSV 失败:{e}"
        out = [f"CSV {p.name} — {len(rows)} 行"]
        for i, r in enumerate(rows[:max_rows]):
            out.append(f"{i + 1} | " + " | ".join(_cell(c) for c in r))
        if len(rows) > max_rows:
            out.append(f"…(共 {len(rows)} 行,超过 {max_rows} 已截断)")
        return "\n".join(out)
    if ext not in (".xlsx", ".xlsm"):
        return f"error: 暂不支持 {ext}(支持 .xlsx/.xlsm/.csv)"
    try:
        import openpyxl

        wb_v = openpyxl.load_workbook(p, data_only=True)   # Excel-cached computed values
        wb_f = openpyxl.load_workbook(p, data_only=False)  # formulas
    except Exception as e:  # noqa: BLE001
        return f"error: 打开工作簿失败:{e}"
    names = [sheet] if sheet else wb_f.sheetnames
    out = [f"工作簿 {p.name} — sheets({len(wb_f.sheetnames)}): {wb_f.sheetnames}"]
    for nm in names:
        if nm not in wb_f.sheetnames:
            out.append(f"\n[sheet「{nm}」不存在]")
            continue
        wf, wv = wb_f[nm], wb_v[nm]
        merged = [str(r) for r in wf.merged_cells.ranges]
        out.append(f"\n## sheet「{nm}」 dims={wf.dimensions} 合并区={merged or '无'}")
        formulas = []
        for i, (rf, rv) in enumerate(zip(wf.iter_rows(), wv.iter_rows())):
            if i >= max_rows:
                out.append(f"…(超过 {max_rows} 行已截断)")
                break
            vals = [_cell(c.value) for c in rv]
            while vals and vals[-1] == "":
                vals.pop()
            if vals:
                out.append(f"{i + 1} | " + " | ".join(vals))
            for c in rf:
                if isinstance(c.value, str) and c.value.startswith("="):
                    formulas.append(f"{c.coordinate}: {c.value}")
        if formulas:
            out.append("公式: " + " ; ".join(formulas[:300]))
    return "\n".join(out)


class SheetSpec(BaseModel):
    name: str
    rows: list[list[str]]
    columns: list[str] | None = None


class SlideSpec(BaseModel):
    title: str
    bullets: list[str] = []
    notes: str | None = None


def _save(title: str, ext: str, data: bytes) -> str:
    path = _filename(title, ext)
    resolve_in_workspace(path).write_bytes(data)
    return path


async def create_document(title: str, content: str, format: str = "docx") -> str:
    """用 Markdown 内容(content)生成一篇文档并存到 workspace,返回文件路径——成功后会自动作为可下载成果展示给用户。format 可选 docx|pdf|html|md。适合产出报告/方案/说明等长文。"""
    fmt = (format or "docx").lower()
    try:
        if fmt == "md":
            return _save(title, "md", content.encode("utf-8"))
        if fmt == "html":
            return _save(title, "html", md_to_html_doc(title, content).encode("utf-8"))
        if fmt == "docx":
            return _save(title, "docx", md_to_docx(title, content))
        if fmt == "pdf":
            return _save(title, "pdf", await md_to_pdf(title, content))
    except Exception as e:  # noqa: BLE001 — surface to the model so it can react/retry
        return f"error: 生成文档失败:{e}"
    return f"error: 不支持的格式 {format}(可选 docx|pdf|html|md)"


def create_spreadsheet(title: str, sheets: list[SheetSpec]) -> str:
    """生成 Excel(.xlsx)并存到 workspace,返回文件路径——成功后自动作为可下载成果展示。每个 sheet 单独一页:name 是页名,rows 是行(每行一组单元格),columns 是可选表头。百分比等格式直接把 "85%" 这类字符串放进单元格。"""
    if not sheets:
        return "error: sheets 不能为空"
    try:
        data = rows_to_xlsx([(s.name, s.columns, s.rows) for s in sheets])
        return _save(title, "xlsx", data)
    except Exception as e:  # noqa: BLE001
        return f"error: 生成 Excel 失败:{e}"


def create_slides(title: str, slides: list[SlideSpec]) -> str:
    """生成 PowerPoint(.pptx)并存到 workspace,返回文件路径——成功后自动作为可下载成果展示。封面用 title;之后每个 slide 一页:title 页标题,bullets 要点(每条一行),notes 可选演讲备注。每页一个清晰要点,别把整段塞进一页。"""
    if not slides:
        return "error: slides 不能为空"
    try:
        outline = [{"title": s.title, "bullets": s.bullets, "notes": s.notes} for s in slides]
        return _save(title, "pptx", outline_to_pptx(title, outline))
    except Exception as e:  # noqa: BLE001
        return f"error: 生成幻灯片失败:{e}"
