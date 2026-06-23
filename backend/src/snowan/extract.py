"""Text extraction from documents, shared by chat attachments and the local-folder
knowledge source. Office/PDF/EPUB parsers are imported lazily."""
import io
import mimetypes
import re
from pathlib import Path

import nh3


def text_of(html: str) -> str:
    """Plain text of an HTML fragment — nh3 strips tags, then collapse whitespace."""
    if not html:
        return ""
    return re.sub(r"\s+", " ", nh3.clean(html, tags=set())).strip()


def extract_text(name: str, mime: str, raw: bytes) -> str | None:
    """Readable text from office/pdf documents. None if not a known doc type."""
    ext = name.lower().rsplit(".", 1)[-1] if "." in name else ""
    try:
        if ext in ("xlsx", "xlsm") or "spreadsheetml" in mime:
            from openpyxl import load_workbook

            wb = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
            out: list[str] = []
            for ws in wb.worksheets:
                out.append(f"## {ws.title}")
                for row in ws.iter_rows(values_only=True):
                    cells = [str(c) for c in row if c is not None]
                    if cells:
                        out.append(" | ".join(cells))
            return "\n".join(out)
        if ext == "docx" or "wordprocessingml" in mime:
            import docx

            d = docx.Document(io.BytesIO(raw))
            return "\n".join(p.text for p in d.paragraphs if p.text.strip())
        if ext == "pdf" or mime == "application/pdf":
            from pypdf import PdfReader

            reader = PdfReader(io.BytesIO(raw))
            return "\n\n".join((pg.extract_text() or "") for pg in reader.pages)
        if ext == "pptx" or "presentationml" in mime:
            from pptx import Presentation

            prs = Presentation(io.BytesIO(raw))
            out2: list[str] = []
            for i, slide in enumerate(prs.slides, 1):
                out2.append(f"## Slide {i}")
                for shape in slide.shapes:
                    if shape.has_text_frame and shape.text_frame.text.strip():
                        out2.append(shape.text_frame.text)
            return "\n".join(out2)
    except Exception:  # noqa: BLE001 — an unparseable doc just falls through
        return None
    return None


def extract_file(path: Path) -> str | None:
    """Text from a file on disk: office/pdf via extract_text, else a utf-8 decode.
    None if unreadable or binary."""
    try:
        raw = path.read_bytes()
    except OSError:
        return None
    mime = mimetypes.guess_type(path.name)[0] or ""
    text = extract_text(path.name, mime, raw)
    if text is not None:
        return text
    try:
        return raw.decode("utf-8")
    except UnicodeDecodeError:
        return None


# --- EPUB ------------------------------------------------------------------

# Same allow-list discipline as reading.sanitize_html: keep semantic structure,
# drop scripts/handlers, harden links. Defined here so the books module gets clean
# chapter HTML without importing the reading DAO.
_EPUB_TAGS = nh3.ALLOWED_TAGS | {
    "figure", "figcaption", "h1", "h2", "h3", "h4", "h5", "h6", "pre", "hr", "span",
}
_EPUB_LINK_REL = "noopener noreferrer nofollow"


def _epub_text(html: str) -> str:
    """Plaintext of a chapter: nh3 strips every tag, then whitespace is collapsed."""
    if not html:
        return ""
    return re.sub(r"\s+", " ", nh3.clean(html, tags=set())).strip()


def _epub_title(html: str) -> str:
    """First heading or non-empty line of a chapter, used when the TOC has no label."""
    m = re.search(r"<h[1-6][^>]*>(.*?)</h[1-6]>", html, re.IGNORECASE | re.DOTALL)
    if m:
        t = _epub_text(m.group(1))
        if t:
            return t[:120]
    for line in _epub_text(html).split(". "):
        line = line.strip()
        if line:
            return line[:120]
    return ""


def _toc_titles(toc) -> dict[str, str]:
    """Flatten an ebooklib TOC (Link / Section / nested tuples) into {href_path: label}.
    The href is split on '#' so a chapter file maps even when the TOC points at an anchor."""
    out: dict[str, str] = {}

    def walk(node) -> None:
        if isinstance(node, (list, tuple)):
            for child in node:
                walk(child)
            return
        href = getattr(node, "href", None)
        title = getattr(node, "title", None)
        if href and title:
            out.setdefault(href.split("#", 1)[0], title.strip())

    walk(toc)
    return out


def _cover_bytes(book) -> bytes | None:
    """Cover image via epub3 (properties='cover-image') or epub2 (<meta name='cover'>),
    falling back to ebooklib's dedicated cover item."""
    from ebooklib import ITEM_COVER, ITEM_IMAGE

    for item in book.get_items_of_type(ITEM_IMAGE):
        props = getattr(item, "properties", None) or []
        if "cover-image" in props:
            return item.get_content()
    meta = book.get_metadata("OPF", "cover")  # epub2: <meta name="cover" content="id">
    if meta:
        cover_id = meta[0][1].get("content")
        item = book.get_item_with_id(cover_id) if cover_id else None
        if item is not None:
            return item.get_content()
    for item in book.get_items_of_type(ITEM_COVER):
        return item.get_content()
    return None


def parse_epub(raw: bytes) -> dict:
    """Parse an EPUB into {title, author, language, cover, chapters}. Each chapter is
    one spine XHTML document in reading order: {num, title, text, html}. `html` is
    nh3-sanitized, `text` is collapsed plaintext, `title` comes from the TOC (falling
    back to the doc's first heading). One unreadable chapter is skipped, not fatal."""
    import tempfile
    import warnings

    from ebooklib import ITEM_DOCUMENT, epub

    # read_epub wants a path on disk; a temp file is the most robust input across
    # ebooklib versions. ignore_ncx avoids a noisy UserWarning on epub3 books.
    with tempfile.NamedTemporaryFile(suffix=".epub") as tmp:
        tmp.write(raw)
        tmp.flush()
        with warnings.catch_warnings():
            warnings.simplefilter("ignore", category=UserWarning)
            book = epub.read_epub(tmp.name, options={"ignore_ncx": True})

    def _meta(name: str) -> str:
        values = book.get_metadata("DC", name)
        return values[0][0].strip() if values and values[0][0] else ""

    toc = _toc_titles(book.toc)
    chapters: list[dict] = []
    num = 0
    for spine_id, _linear in book.spine:
        item = book.get_item_with_id(spine_id)
        if item is None or item.get_type() != ITEM_DOCUMENT:
            continue
        # The epub3 nav/TOC document is structure, not a readable chapter. ebooklib
        # keeps it as an EpubNav instance; third-party books tag it via OPF properties.
        if isinstance(item, epub.EpubNav) or "nav" in (getattr(item, "properties", None) or []):
            continue
        try:
            raw_html = item.get_content().decode("utf-8", "replace")
        except Exception:  # noqa: BLE001 — one bad chapter must not abort the import
            continue
        text = _epub_text(raw_html)
        if not text:
            continue  # skip nav/cover XHTML and other empty documents
        href = (item.get_name() or "").split("#", 1)[0]
        title = toc.get(href) or _epub_title(raw_html) or f"第 {num + 1} 章"
        chapters.append({
            "num": num,
            "title": title,
            "text": text,
            "html": nh3.clean(raw_html, tags=_EPUB_TAGS, link_rel=_EPUB_LINK_REL),
        })
        num += 1

    return {
        "title": _meta("title") or book.title or "未命名书籍",
        "author": _meta("creator"),
        "language": _meta("language") or (book.language or ""),
        "cover": _cover_bytes(book),
        "chapters": chapters,
    }
